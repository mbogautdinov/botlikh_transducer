"""Botlikh corpus -> one row per word occurrence.
Usage (from the root of the repository):  python3 scripts/corpus_words.py [--xlsx PATH] [--out PATH]
Default input: evaluation/corpus/botlikh_corpus_merged.xlsx, default output: evaluation/corpus/botlikh_corpus_words.xlsx

Words are taken from {{...}} in 'Botlikh (raw markup)'; segmentation and gloss are aligned by the word order (punctuation tokens skipped).
Words with a RUS mark in the gloss, then words with an AVAR mark, go to the end. Part of speech is left empty (filled in by hand).
The table in the repository has the parts of speech filled in by hand, so an existing file is overwritten only with --force.
"""
import argparse
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CORPUS = ROOT / "evaluation" / "corpus"

PUNCT = re.compile(r'[-–—,.;:!?"«»()…]+')
LATIN = re.compile(r'[A-Za-z]')
CYR = re.compile(r'[А-Яа-яЁё]')

GLOSS_OVERRIDE = {60: ["an.pl-hello", "from morning // с утра", "everyone // все"]}   # a gloss without '//' in an Avar sentence
SEG_MERGE = {56: (8, 9), 256: (4, 5)}      # 'то есть' is one word in the text and two in the segmentation
WORDS_OVERRIDE = {195: ["Дий", "на", "хъуча", "баила", "иде"]}   # broken markup '{{на {{хъуча}}'
DUP_GLOSS = {246: 3}                       # one gloss for two words ('лъугьун руго')


def avar_groups(gloss):
    toks = gloss.replace("//", " // ").replace("<AVAR>", " <AVAR> ").split()
    groups, en, ru, state = [], [], [], "en"

    def flush():
        nonlocal en, ru, state
        if en or ru:
            groups.append(" ".join(en) + (" // " + " ".join(ru) if ru else ""))
        en, ru, state = [], [], "en"
    for t in toks:
        if t == "//":
            state = "ru"
        elif t == "<AVAR>":
            flush()
        elif state == "ru" and LATIN.search(t) and not CYR.search(t):
            flush()
            en.append(t)
        else:
            (ru if state == "ru" else en).append(t)
    flush()
    return groups


def kind(gloss):
    if "(AVAR" in gloss or "<AVAR" in gloss or "//" in gloss:
        return "AVAR"
    if "RUS" in gloss:
        return "RUS"
    return "Botlikh"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default=str(CORPUS / "botlikh_corpus_merged.xlsx"))
    ap.add_argument("--out", default=str(CORPUS / "botlikh_corpus_words.xlsx"))
    ap.add_argument("--force", action="store_true", help="overwrite the output file (the parts of speech filled in by hand are lost)")
    a = ap.parse_args()
    if Path(a.out).exists() and not a.force:
        sys.exit(f"{a.out} exists; give another --out or use --force")
    d = pd.read_excel(a.xlsx)
    rows, problems = [], []
    for i, r in d.iterrows():
        raw = str(r["Botlikh (raw markup)"])
        words = WORDS_OVERRIDE.get(i) or [w for w in re.findall(r"\{\{(.*?)\}\}", raw) if not PUNCT.fullmatch(w)]
        seg = [t for t in (str(r["Segmentation"]) if pd.notna(r["Segmentation"]) else "").split() if not PUNCT.fullmatch(t)]
        if i in SEG_MERGE:
            x, y = SEG_MERGE[i]
            seg[x - 1:y] = [" ".join(seg[x - 1:y])]
        g = str(r["Gloss"]) if pd.notna(r["Gloss"]) else ""
        gl = GLOSS_OVERRIDE.get(i) or (avar_groups(g) if "//" in g else g.split())
        notes = [[] for _ in words]
        if i in DUP_GLOSS:
            gl.insert(DUP_GLOSS[i] + 1, gl[DUP_GLOSS[i]])
            notes[DUP_GLOSS[i]].append("одна глосса на два слова")
            notes[DUP_GLOSS[i] + 1].append("одна глосса на два слова")
        if i in SEG_MERGE:
            notes[SEG_MERGE[i][0] - 1].append("в сегментации два слова")
        if not (len(words) == len(seg) == len(gl)):
            problems.append((i, len(words), len(seg), len(gl)))
            for n in notes:
                n.append("выравнивание не сошлось, проверить")
        for j, w in enumerate(words):
            gj = gl[j] if j < len(gl) else ""
            rows.append({"Язык слова": kind(gj), "Слово": w, "Сегментация": seg[j] if j < len(seg) else "", "Глосса": gj,
                         "Часть речи": "", "Предложение (ботлихский)": r["Botlikh"], "Предложение (русский)": r["Russian"],
                         "Предложение (английский)": r["English"], "Документ": r["Document"], "Говорящий": r["Speaker"],
                         "№ предложения": i + 1, "№ слова": j + 1, "Примечание": "; ".join(notes[j])})
    out = pd.DataFrame(rows)
    order = {"Botlikh": 0, "RUS": 1, "AVAR": 2}
    out = out.sort_values(by=["Язык слова", "№ предложения", "№ слова"], key=lambda c: c.map(order) if c.name == "Язык слова" else c, kind="stable")
    with pd.ExcelWriter(a.out, engine="openpyxl") as w:
        out.to_excel(w, index=False, sheet_name="Words")
        ws = w.sheets["Words"]
        for k, col in enumerate(out.columns, 1):
            ws.column_dimensions[get_column_letter(k)].width = {"Предложение (ботлихский)": 45, "Предложение (русский)": 45, "Предложение (английский)": 45,
                                                               "Документ": 30, "Глосса": 30, "Сегментация": 22}.get(col, 16)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDDDDD")
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
    print(len(out), "occurrences;", out["Язык слова"].value_counts().to_dict(), "; not aligned:", problems)


if __name__ == "__main__":
    main()
